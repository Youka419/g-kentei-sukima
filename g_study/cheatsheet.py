from __future__ import annotations

from datetime import datetime
from io import BytesIO
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from .config import DATA_DIR, EXCEL_COLUMNS, EXCEL_PATH, WORD_PATH
from .runtime import disk_writable
from .terms import SEED_CHEATSHEET_TERMS, find_term


def ensure_data_dir() -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)


def empty_frame() -> pd.DataFrame:
    return pd.DataFrame(columns=EXCEL_COLUMNS)


def _row_from_term(term: dict, extra: dict | None = None) -> dict:
    extra = extra or {}
    row = {
        "追加日": extra.get("追加日", datetime.now().strftime("%Y-%m-%d %H:%M")),
        "分野": term.get("field", extra.get("分野", "")),
        "章": term.get("chapter", extra.get("章", "")),
        "節": term.get("section", extra.get("節", "")),
        "用語": extra.get("用語", term.get("term", "")),
        "定義": extra.get("定義", term.get("definition", "")),
        "試験ポイント": extra.get("試験ポイント", term.get("exam_point", "")),
        "問題文": extra.get("問題文", ""),
        "自分の解答": extra.get("自分の解答", ""),
        "正解": extra.get("正解", ""),
        "メモ": extra.get("メモ", "初期チートシートから転記"),
        "復習フラグ": extra.get("復習フラグ", "要復習"),
        "出典PDF": extra.get("出典PDF", term.get("pdf", "")),
    }
    return {k: row.get(k, "") for k in EXCEL_COLUMNS}


def seed_from_existing_cheatsheet() -> pd.DataFrame:
    rows = []
    for name in SEED_CHEATSHEET_TERMS:
        term = find_term(name)
        if term:
            rows.append(_row_from_term(term, {"メモ": "既存PDFチートシートの重要語", "復習フラグ": "重要"}))
    return pd.DataFrame(rows, columns=EXCEL_COLUMNS)


def normalize(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    for col in EXCEL_COLUMNS:
        if col not in out.columns:
            out[col] = ""
    return out[EXCEL_COLUMNS].fillna("")


def load_excel(path: Path | None = None) -> pd.DataFrame:
    target = path or EXCEL_PATH
    if target.exists():
        df = pd.read_excel(target, engine="openpyxl")
        return normalize(df)
    df = seed_from_existing_cheatsheet()
    if disk_writable(DATA_DIR):
        save_excel(df, target)
        export_word(df, WORD_PATH)
    return df


def load_excel_bytes(data: bytes) -> pd.DataFrame:
    df = pd.read_excel(BytesIO(data), engine="openpyxl")
    return normalize(df)


def excel_bytes(df: pd.DataFrame) -> bytes:
    buf = BytesIO()
    out = df.copy() if "不正解回数" in df.columns else normalize(df)
    out.fillna("").to_excel(buf, index=False, engine="openpyxl")
    return buf.getvalue()


def save_excel(df: pd.DataFrame, path: Path | None = None) -> Path:
    ensure_data_dir()
    target = path or EXCEL_PATH
    normalize(df).to_excel(target, index=False, engine="openpyxl")
    return target


def add_entry(df: pd.DataFrame, payload: dict) -> pd.DataFrame:
    term = find_term(payload.get("用語", "")) or {}
    row = _row_from_term(term, payload)
    added = pd.concat([df, pd.DataFrame([row])], ignore_index=True)
    return added


def already_has(df: pd.DataFrame, term: str, question: str = "") -> bool:
    if df.empty:
        return False
    same_term = df["用語"].astype(str) == str(term)
    if question:
        return bool((same_term & (df["問題文"].astype(str) == str(question))).any())
    return bool(same_term.any())


def chronic_miss_frame(sheet: pd.DataFrame, progress: dict, min_count: int = 10) -> pd.DataFrame:
    from .progress import chronic_misses

    cols = [
        "不正解回数",
        "分野",
        "章",
        "節",
        "用語",
        "定義",
        "試験ポイント",
        "問題文",
        "自分の解答",
        "正解",
        "メモ",
    ]
    rows = []
    for item in chronic_misses(progress, min_count):
        extra = {"問題文": "", "自分の解答": "", "正解": "", "メモ": ""}
        if sheet is not None and not sheet.empty and "用語" in sheet.columns:
            hit = sheet[sheet["用語"].astype(str) == str(item["用語"])]
            if not hit.empty:
                last = hit.iloc[-1]
                extra = {
                    "問題文": last.get("問題文", ""),
                    "自分の解答": last.get("自分の解答", ""),
                    "正解": last.get("正解", ""),
                    "メモ": last.get("メモ", ""),
                }
        rows.append({**item, **extra})
    if not rows:
        return pd.DataFrame(columns=cols)
    return pd.DataFrame(rows)[cols].fillna("")


def _set_run_font(run, size: int = 11, bold: bool = False, color: tuple[int, int, int] | None = None) -> None:
    run.font.size = Pt(size)
    run.bold = bold
    run.font.name = "Yu Gothic"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "游ゴシック")
    if color:
        run.font.color.rgb = RGBColor(*color)


def _build_word(df: pd.DataFrame) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Pt(54)
    section.left_margin = section.right_margin = Pt(54)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(
        "G検定 10回以上間違えた問題" if "不正解回数" in df.columns else "G検定 弱点チートシート"
    )
    _set_run_font(run, 20, bold=True, color=(15, 61, 107))

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run(
        f"更新: {datetime.now().strftime('%Y-%m-%d %H:%M')}　／　件数: {len(df)}　／　Excelから自動生成"
    )
    _set_run_font(run, 9, color=(90, 90, 90))

    note = doc.add_paragraph()
    run = note.add_run(
        "間違えた問題や覚えたい用語をExcelで管理し、このWordに書き出しています。"
        "公式シラバス（G2024#6〜）の章立てで並べています。"
    )
    _set_run_font(run, 10)

    if df.empty:
        p = doc.add_paragraph()
        run = p.add_run("まだ項目がありません。アプリのクイズや手動追加から登録してください。")
        _set_run_font(run, 11)
        return doc

    grouped = df.fillna("")
    for chapter, g1 in grouped.groupby("章", sort=False):
        heading = doc.add_paragraph()
        run = heading.add_run(str(chapter) if chapter else "未分類")
        _set_run_font(run, 14, bold=True, color=(15, 61, 107))
        for _, row in g1.iterrows():
            p = doc.add_paragraph()
            run = p.add_run(f"■ {row['用語']}")
            _set_run_font(run, 12, bold=True)
            if row.get("節"):
                run = p.add_run(f"　[{row['節']}]")
                _set_run_font(run, 9, color=(100, 100, 100))

            if row.get("不正解回数") not in ("", None):
                d = doc.add_paragraph()
                run = d.add_run(f"不正解回数: {row['不正解回数']}")
                _set_run_font(run, 11, bold=True, color=(153, 0, 0))
            if row.get("定義"):
                d = doc.add_paragraph()
                run = d.add_run(f"定義: {row['定義']}")
                _set_run_font(run, 11)
            if row.get("試験ポイント"):
                d = doc.add_paragraph()
                run = d.add_run(f"試験: {row['試験ポイント']}")
                _set_run_font(run, 10, color=(40, 40, 40))
            if row.get("問題文"):
                d = doc.add_paragraph()
                run = d.add_run(f"問題: {row['問題文']}")
                _set_run_font(run, 10)
            if row.get("自分の解答") or row.get("正解"):
                d = doc.add_paragraph()
                run = d.add_run(f"自分の解答: {row.get('自分の解答', '')}　／　正解: {row.get('正解', '')}")
                _set_run_font(run, 10, color=(153, 0, 0))
            if row.get("メモ"):
                d = doc.add_paragraph()
                run = d.add_run(f"メモ: {row['メモ']}")
                _set_run_font(run, 10)
            flag = row.get("復習フラグ", "")
            if flag:
                d = doc.add_paragraph()
                run = d.add_run(f"フラグ: {flag}")
                _set_run_font(run, 9, color=(15, 61, 107))
    return doc


def export_word(df: pd.DataFrame, path: Path | None = None) -> Path:
    ensure_data_dir()
    target = path or WORD_PATH
    _build_word(df).save(target)
    return target


def word_bytes(df: pd.DataFrame) -> bytes:
    buf = BytesIO()
    _build_word(df).save(buf)
    return buf.getvalue()


def save_both(df: pd.DataFrame) -> tuple[Path | None, Path | None]:
    if not disk_writable(DATA_DIR):
        return None, None
    return save_excel(df), export_word(df)
